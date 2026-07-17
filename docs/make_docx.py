#!/usr/bin/env python3
"""Generate the shareable Word doc (Content_Freshness_Radar.docx) for the
hackathon organizing team. Run: python3 docs/make_docx.py"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "dashboard.png")
OUT = os.path.join(HERE, "Content_Freshness_Radar.docx")

WIKI_BLUE = RGBColor(0x33, 0x66, 0xCC)
INK = RGBColor(0x20, 0x21, 0x22)
SUBTLE = RGBColor(0x54, 0x59, 0x5D)

doc = Document()

# Base font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = INK


def heading(text, size, color=INK, space_before=14, space_after=6, bold=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def body(text, italic=False, color=INK, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


# ---- Title block ----
heading("Content Freshness Radar", 26, WIKI_BLUE, space_before=0, space_after=2)
body("Making Wikipedia's obsolescence visible — and easy to fix.",
     italic=True, color=SUBTLE, size=12, space_after=2)
body("Hackathon submission overview", color=SUBTLE, size=10, space_after=10)

# ---- Challenge ----
heading("The challenge", 15, WIKI_BLUE)
body("Identify and flag outdated articles, obsolete illustrations and "
     "screenshots, and outdated graphics — to encourage the community to "
     "update them. An article may be factually accurate when written but become "
     "misleading over time: a statistic from 2012, an org chart of a "
     "restructured institution, or a political map from before an election. "
     "This obsolescence is usually invisible to readers, and those who notice a "
     "problem rarely have an easy way to act.")

# ---- Solution ----
heading("Our solution", 15, WIKI_BLUE)
body("Content Freshness Radar scans a topic area of a wiki and produces a "
     "prioritized worklist of pages that show real evidence of being outdated, "
     "each with the specific reason and a one-click path to fix it. It is built "
     "to look and feel native to the wiki, using the Wikimedia Codex design "
     "system.")

# ---- Screenshot ----
if os.path.exists(IMG):
    doc.add_picture(IMG, width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = body("The dashboard: an evidence-based flag count, a descriptive age "
               "lens, and an evidence-gated \u201cReview & update\u201d button.",
               italic=True, color=SUBTLE, size=9, space_after=10)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ---- Key decision ----
heading("The key design decision: age is a lens, not a verdict", 15, WIKI_BLUE)
body("Flagging anything not edited in N years is wrong: a history or maths "
     "article can be untouched for years because it is complete, and telling "
     "editors to \u201cupdate\u201d it erodes trust. So we separate freshness "
     "(descriptive age color-coding, a browsing lens) from update evidence "
     "(concrete, page-specific proof of staleness). Only evidence \u2014 the "
     "article's own {{As of|YYYY}} dated statements being old, or an existing "
     "community maintenance tag \u2014 triggers a recommendation. A page edited "
     "yesterday can be flagged if it cites \u201cas of 2012\u201d data; a stable "
     "page untouched for years is left alone.")

# ---- Mapping table ----
heading("How it addresses the challenge", 15, WIKI_BLUE)
rows = [
    ("Challenge example", "How we do it"),
    ("Color-code by how long since last update",
     "Neutral GREEN / YELLOW / ORANGE / RED age bands from the last substantive edit"),
    ("Flag articles that are outdated",
     "Evidence-based: {{As of|YYYY}} dated statements + community tags, not raw edit age"),
    ("Identify outdated data in charts / statistics",
     "Surfaces the year a figure is marked \u201cas of\u201d (e.g. \u201ccites 'as of 2012' data \u2014 14 years old\u201d)"),
    ("Encourage improvement of outdated content",
     "Flagged rows get a \u201cReview & update\u201d button; {{Update}} suggested only with evidence and if not already tagged"),
    ("Identify outdated screenshots / images",
     "Deliberately secondary/experimental \u2014 upload age is a weak signal; most images never need updating"),
]
table = doc.add_table(rows=0, cols=2)
table.style = "Light Grid Accent 1"
for i, (a, b) in enumerate(rows):
    cells = table.add_row().cells
    cells[0].text = a
    cells[1].text = b
    if i == 0:
        for c in cells:
            for para in c.paragraphs:
                for run in para.runs:
                    run.bold = True
body("", space_after=4)

# ---- Precision ----
heading("What makes it accurate (not just noisy)", 15, WIKI_BLUE)
body("The hard part is precision — flagging stable content (a maths proof, a "
     "historical battle) erodes community trust. Two design choices keep the "
     "signal clean:")
bullet("A page's freshness clock is only reset by a real content edit. Minor "
       "edits, bot edits, reverts, and cosmetic tweaks are filtered out.",
       bold_prefix="\u201cSubstantive edit\u201d detection. ")
bullet("Old age never triggers a recommendation on its own. Only concrete, "
       "page-specific signals ({{As of}} dated statements being old, or an "
       "existing community update tag) do. Topic volatility is used only to "
       "rank candidates, never to assert a page is outdated.",
       bold_prefix="Evidence-gated recommendations. ")

# ---- How it works ----
heading("How it works", 15, WIKI_BLUE)
bullet("Live MediaWiki Action API + Pageviews API for revisions, pageviews, and "
       "categories, with a policy-compliant User-Agent.", bold_prefix="Data: ")
bullet("Python, standard library only — no dependencies to install.",
       bold_prefix="Backend: ")
bullet("Vue 3 + Wikimedia Codex, loaded via an ES-module import map — no build "
       "step and no npm required.", bold_prefix="Frontend: ")
bullet("The same logic maps to the Wikimedia replica databases / dumps for "
       "whole-wiki analysis (grounded in the real MediaWiki 1.45 schema).",
       bold_prefix="Scale path: ")

# ---- Results ----
heading("Live results", 15, WIKI_BLUE)
body("Run on Demographics of Asia (30 pages), the tool flags exactly 3 with "
     "evidence — e.g. Demographics of Thailand, edited days ago but still citing "
     "\u201cas of 2016\u201d data — while leaving 27 stable pages untouched. On "
     "Economy of India it flags a statistic still marked \u201cas of 2012.\u201d")

# ---- Try it ----
heading("Try it in two commands", 15, WIKI_BLUE)
for line in [
    "python3 run_pipeline.py --category \"Demographics of Asia\" --recursive --limit 30",
    "cd webapp && python3 -m http.server 8777   # open http://localhost:8777",
]:
    p = doc.add_paragraph()
    r = p.add_run(line)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
body("Point it at any category on any MediaWiki wiki with --wiki and --category.",
     color=SUBTLE, size=10)

# ---- Roadmap ----
heading("Status & roadmap", 15, WIKI_BLUE)
bullet("neutral freshness bands, substantive-edit filtering, evidence-based "
       "update recommendation ({{As of}} + maintenance tags), pageview "
       "priority, recursive category scanning, and the Codex dashboard with an "
       "evidence-gated \u201cReview & update\u201d CTA.",
       bold_prefix="Working today: ")
bullet("a Wikidata ground-truth check (compare a stated figure to Wikidata's "
       "current value) as a stronger evidence source, parse data-years from "
       "tables/charts, and replica-DB/dump ingestion for full-wiki coverage.",
       bold_prefix="Next: ")

doc.save(OUT)
print("Wrote", OUT)
